import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

REPO_ROOT = Path(__file__).resolve().parents[4]
REPO_TMP = REPO_ROOT / "tmp"


class SyntheticDocxRoundtripTests(unittest.TestCase):
    def test_generated_docx_reopens_with_expected_semantic_structure(self):
        self.assertTrue(REPO_TMP.resolve().is_relative_to(REPO_ROOT.resolve()))
        REPO_TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix="word-docx-smoke-", dir=REPO_TMP) as temp:
            path = Path(temp) / "synthetic.docx"
            doc = Document()
            doc.add_heading("Synthetic handoff", level=1)
            doc.add_heading("Scope", level=2)
            doc.add_paragraph("This is synthetic content for a structural smoke test.")
            doc.add_paragraph("First check", style="List Bullet")
            doc.save(path)

            self.assertTrue(path.is_file())
            with zipfile.ZipFile(path) as archive:
                self.assertIsNone(archive.testzip())
                self.assertIn("word/document.xml", archive.namelist())

            reopened = Document(path)
            paragraphs = [p for p in reopened.paragraphs if p.text]
            self.assertEqual(paragraphs[0].style.name, "Heading 1")
            self.assertEqual(paragraphs[1].style.name, "Heading 2")
            self.assertEqual(paragraphs[3].style.name, "List Bullet")
            self.assertIn("synthetic content", " ".join(p.text for p in paragraphs))




if __name__ == "__main__":
    unittest.main()
