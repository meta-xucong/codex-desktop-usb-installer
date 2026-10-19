from __future__ import annotations

import re
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl import Workbook, load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]


class OpenpyxlFormatContractTests(unittest.TestCase):
    def test_actual_bundled_openpyxl_version(self):
        self.assertEqual(openpyxl.__version__, "3.1.5")

    def test_xlsx_roundtrip_and_formula_cache_is_not_calculated(self):
        with tempfile.TemporaryDirectory(prefix="excel-xlsx-contract-") as directory:
            path = Path(directory) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "read-write"
            sheet["A2"] = "=1+2"
            workbook.save(path)

            formulas = load_workbook(path, data_only=False)
            self.assertEqual(formulas.active["A1"].value, "read-write")
            self.assertEqual(formulas.active["A2"].value, "=1+2")
            cached = load_workbook(path, data_only=True)
            self.assertIsNone(cached.active["A2"].value)

            formulas.active["A3"] = "written"
            formulas.save(path)
            self.assertEqual(load_workbook(path, data_only=False).active["A3"].value, "written")

    def test_macro_free_xlsm_keep_vba_roundtrip(self):
        with tempfile.TemporaryDirectory(prefix="excel-xlsm-contract-") as directory:
            path = Path(directory) / "sample.xlsm"
            workbook = Workbook()
            workbook.active["A1"] = "macro-free xlsm package"
            workbook.save(path)

            loaded = load_workbook(path, keep_vba=True)
            self.assertEqual(loaded.active["A1"].value, "macro-free xlsm package")
            self.assertIsNotNone(loaded.vba_archive)
            loaded.active["A2"] = "roundtrip"
            loaded.save(path)
            self.assertEqual(load_workbook(path, keep_vba=True).active["A2"].value, "roundtrip")

    def test_docs_limit_formats_and_formula_claims(self):
        skill = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        setup = (SKILL_DIR / "setup.md").read_text(encoding="utf-8")
        when_to_use = re.search(r"## When to Use(.*?)(?=\n## |\Z)", skill, re.S)
        self.assertIsNotNone(when_to_use)
        self.assertNotRegex(when_to_use.group(1), r"\.xls(?:,|\s|\))")
        self.assertIn("does not read or write legacy `.xls`", skill)
        self.assertIn("keep_vba=True", skill)
        self.assertIn("does not calculate formulas", skill)
        self.assertIn("macros are preserved, never run by `openpyxl`", skill)
        self.assertIn("Excel or LibreOffice", setup)
        self.assertIn("does not run macros or calculate formulas", setup)
        self.assertNotIn("formula evaluation", skill.lower())
        self.assertNotIn("formula evaluation", setup.lower())


if __name__ == "__main__":
    unittest.main()
